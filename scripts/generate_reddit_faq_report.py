"""Builds a Reddit FAQ DOCX from Reddit Research results - real
question-shaped post titles found on a subreddit, with where each one was
found. Deliberately minimal: just "Qn. <question>" + "Found on" per
question. An "Answer" row is added only when a real AI provider key is
configured and a draft was actually generated from the thread's top
comments - if no AI key works, the report stops at Question + Found on,
nothing else.
"""
import json
import re
import urllib.error
import urllib.request

from docx import Document
from docx.shared import Inches, Pt, RGBColor

_QUESTION_STARTERS = (
    "what", "why", "how", "is", "are", "does", "do", "can", "could",
    "should", "will", "when", "where", "who", "which", "would",
)

COL0_WIDTH = Inches(1.4)
COL1_WIDTH = Inches(5.1)


def _looks_like_question(title):
    t = title.strip()
    if not t:
        return False
    if t.rstrip().endswith("?"):
        return True
    first_word = re.split(r"[\s:]+", t, 1)[0].lower().strip(".,!")
    return first_word in _QUESTION_STARTERS


def select_questions(results, max_questions=15):
    """Pick question-shaped posts from Reddit Research results, highest
    engagement (score) first, skipping near-duplicate titles (the same
    question cross-posted/re-asked across threads)."""
    candidates = [r for r in results if _looks_like_question(r["post"]["title"])]
    candidates.sort(key=lambda r: int(r["post"].get("score") or 0), reverse=True)
    seen_titles = set()
    deduped = []
    for r in candidates:
        norm = re.sub(r"[^a-z0-9 ]", "", r["post"]["title"].lower()).strip()
        if norm in seen_titles:
            continue
        seen_titles.add(norm)
        deduped.append(r)
    return deduped[:max_questions]


def _found_on_text(post):
    return (f'r/{post["subreddit"]} - "{post["title"]}" '
            f'({post.get("score", 0)} upvotes, {post.get("num_comments", 0)} comments)')


# --------------------------------------------------------------------------- #
# AI answer drafting - same free-tier-first/paid-last provider chain used by
# the On-Page report (Gemini free -> Groq free -> OpenRouter free -> OpenAI
# paid), re-implemented standalone here since that module is written for
# structured on-page JSON suggestions, not free-form FAQ answers, and runs
# as a subprocess rather than being imported in-process like this is.
# --------------------------------------------------------------------------- #
def _call_gemini(prompt, key, model="gemini-2.5-flash"):
    url = (f"https://generativelanguage.googleapis.com/v1beta/models/"
           f"{model}:generateContent?key={key}")
    body = json.dumps({"contents": [{"parts": [{"text": prompt}]}],
                        "generationConfig": {"temperature": 0.5}}).encode("utf-8")
    req = urllib.request.Request(url, data=body, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as r:
        data = json.loads(r.read())
    return data["candidates"][0]["content"]["parts"][0]["text"].strip()


def _call_openai_compatible(prompt, key, url, model):
    payload = {"model": model, "messages": [{"role": "user", "content": prompt}], "temperature": 0.5}
    req = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers={
        "Content-Type": "application/json", "Authorization": f"Bearer {key}",
    })
    with urllib.request.urlopen(req, timeout=30) as r:
        data = json.loads(r.read())
    return data["choices"][0]["message"]["content"].strip()


def _draft_answer(question, comments, api_keys):
    """Tries each configured provider in order, free tiers first; returns a
    drafted answer string, or None if no provider is configured/succeeds."""
    comment_text = "\n".join(f'- {c.get("body", "")}' for c in (comments or [])[:5]) or "(no comments)"
    prompt = (
        f'A real question was asked on Reddit: "{question}"\n\n'
        f'Real replies from that thread:\n{comment_text}\n\n'
        f'Write a short, direct answer (2-4 sentences, plain English, no Reddit slang, '
        f'no mention of Reddit) suitable for publishing on a company FAQ page, based on '
        f'the genuine information in those replies. If the replies do not actually answer '
        f'the question, reply with exactly: NO_ANSWER'
    )
    providers = []
    if api_keys.get("gemini"):
        providers.append(lambda: _call_gemini(prompt, api_keys["gemini"]))
    if api_keys.get("groq"):
        providers.append(lambda: _call_openai_compatible(
            prompt, api_keys["groq"], "https://api.groq.com/openai/v1/chat/completions",
            "llama-3.3-70b-versatile"))
    if api_keys.get("openrouter"):
        providers.append(lambda: _call_openai_compatible(
            prompt, api_keys["openrouter"], "https://openrouter.ai/api/v1/chat/completions",
            "meta-llama/llama-3.3-70b-instruct:free"))
    if api_keys.get("openai"):
        providers.append(lambda: _call_openai_compatible(
            prompt, api_keys["openai"], "https://api.openai.com/v1/chat/completions", "gpt-4o-mini"))

    for call in providers:
        try:
            answer = call()
            if answer and "NO_ANSWER" not in answer.upper():
                return answer
        except Exception:
            continue
    return None


# --------------------------------------------------------------------------- #
# Document
# --------------------------------------------------------------------------- #
def _add_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def _add_kv_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Normal Table"
    table.autofit = False
    for ri, (label, value) in enumerate(rows):
        label_cell, value_cell = table.rows[ri].cells
        label_cell.width = COL0_WIDTH
        value_cell.width = COL1_WIDTH
        lr = label_cell.paragraphs[0].add_run(label)
        lr.bold = True
        lr.font.color.rgb = RGBColor(0, 0, 0)
        vr = value_cell.paragraphs[0].add_run(value)
        vr.font.color.rgb = RGBColor(0, 0, 0)
    return table


def build_faq_docx(results, out_path, subreddit, keyword=None, max_questions=15, api_keys=None):
    """results: the list Reddit Research's /api/reddit/search returns
    ([{"post": {...}, "comments": [...]}, ...]). api_keys: optional dict of
    {"gemini"/"groq"/"openrouter"/"openai": key} - when at least one is set
    and succeeds, an "Answer" row is added per question from a real AI draft
    grounded in that thread's actual comments; otherwise each question is
    just "Found on", nothing else."""
    selected = select_questions(results, max_questions=max_questions)
    api_keys = api_keys or {}

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(1)

    for i, item in enumerate(selected, 1):
        post = item["post"]
        _add_heading(doc, f"Q{i}. {post['title']}")
        rows = [("Found on", _found_on_text(post))]
        answer = _draft_answer(post["title"], item.get("comments"), api_keys) if api_keys else None
        if answer:
            rows.append(("Answer", answer))
        _add_kv_table(doc, rows)
        doc.add_paragraph()

    if not selected:
        p = doc.add_paragraph()
        p.add_run("No clearly question-shaped posts were found in this search - try a "
                  "broader time range, a higher post count, or a different subreddit.")

    doc.save(out_path)
    return out_path
